"""Build docs/business_case/TailSignal_Business_Case.xlsx (formulas throughout; recalc with LibreOffice after)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

OUT = Path("docs/business_case/TailSignal_Business_Case.xlsx")
F = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="E8EDF3")
thin = Side(style="thin", color="B7C3D0")
USD = '$#,##0;($#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
YEARS = ["Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]

wb = Workbook()


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)


def header(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=t)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)


def section(ws, row, text, cols=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(bold=True, size=11)
    for j in range(1, cols + 1):
        ws.cell(row=row, column=j).fill = SUB


# ----------------------------------------------------------------------------- Read Me
rm = wb.active
rm.title = "Read Me"
rm.column_dimensions["A"].width = 110
lines = [
    ("TailSignal business case: five-year model", font(bold=True, size=14)),
    ("Purpose: size the four monetization models and show when each product unlocks as the partner network grows.", font()),
    ("", font()),
    ("How to use", font(bold=True, size=11)),
    ("1. Edit only blue cells on the Assumptions sheet. Yellow cells are the key levers.", font()),
    ("2. Every other number is a formula. Results for all three scenarios appear on the Summary sheet.", font()),
    ("3. Scale gates (when drug-safety studies become possible) come from the TailSignal analyses, not from guesses.", font()),
    ("", font()),
    ("Legend", font(bold=True, size=11)),
    ("Blue text = input you can change. Black text = formula. Green text = link to another sheet. Yellow fill = key assumption.", font()),
    ("", font()),
    ("Important", font(bold=True, size=11)),
    ("Prices, adoption rates, headcount and partner ramp are ILLUSTRATIVE ASSUMPTIONS for a portfolio project. Replace them with market quotes and pipeline data before any decision.", font(color="C00000", bold=True)),
    ("Market-context figures are cited on the Assumptions sheet with their sources.", font()),
]
for i, (t, f) in enumerate(lines, 1):
    c = rm.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ----------------------------------------------------------------------------- Assumptions
a = wb.create_sheet("Assumptions")
a.column_dimensions["A"].width = 52
for col in "BCDEF":
    a.column_dimensions[col].width = 14
a.column_dimensions["G"].width = 70
A = {}  # name -> absolute ref


def inp(row, label, value, fmt, name, note=None, key=False, src=None):
    a.cell(row=row, column=1, value=label).font = font()
    c = a.cell(row=row, column=2, value=value)
    c.font = font(BLUE)
    c.number_format = fmt
    if key:
        c.fill = YELLOW
    if note:
        a.cell(row=row, column=7, value=note).font = font(italic=True, color="555555")
    if src:
        c.comment = Comment(src, "TailSignal")
    A[name] = f"Assumptions!$B${row}"


def inp_years(row, label, values, fmt, name, note=None, key=False):
    a.cell(row=row, column=1, value=label).font = font()
    for j, v in enumerate(values):
        c = a.cell(row=row, column=2 + j, value=v)
        c.font = font(BLUE)
        c.number_format = fmt
        if key:
            c.fill = YELLOW
    if note:
        a.cell(row=row, column=7, value=note).font = font(italic=True, color="555555")
    A[name] = row


r = 1
a.cell(row=r, column=1, value="Assumptions").font = font(bold=True, size=14)
r = 3
section(a, r, "Market context (sourced; not used in formulas)", 7)
r += 1
header(a, r, ["Item", "Value", "", "", "", "", "Source"])
ctx = [("US veterinary practices", 34000, NUM, "AVMA 2025 Economic State of the Profession, via dvm360 (2022 count): https://www.dvm360.com/view/2025-economic-state-of-the-veterinary-profession-trends-and-opportunities-for-your-practice"),
       ("Share of practices owned by corporate groups", 0.30, PCT, "Same source: corporate groups own ~30% of US practices and over half of companion-animal revenue"),
       ("US households owning a pet (millions)", 95, NUM, "APPA, 2025: https://americanpetproducts.org/news/u.s.-pet-industry-reaches-158-billion-in-2025-poised-for-continued-growth-in-2026"),
       ("US pet industry spending, 2025 ($bn)", 158, NUM, "APPA, 2025 (same link)")]
for lbl, v, fmt, src in ctx:
    r += 1
    a.cell(row=r, column=1, value=lbl).font = font()
    c = a.cell(row=r, column=2, value=v)
    c.font = font(BLUE)
    c.number_format = fmt
    a.cell(row=r, column=7, value=src).font = font(italic=True, color="555555")
    a.cell(row=r, column=7).alignment = Alignment(wrap_text=True)
r += 2
section(a, r, "Scenarios", 7)
r += 1
header(a, r, ["Scenario", "Partner ramp ×", "Price ×", "", "", "", "Note"])
scen = [("Low", 0.6, 0.8), ("Base", 1.0, 1.0), ("High", 1.3, 1.15)]
A["scen_rows"] = []
for name, rm_, pm in scen:
    r += 1
    a.cell(row=r, column=1, value=name).font = font(bold=True)
    for j, v in enumerate([rm_, pm]):
        c = a.cell(row=r, column=2 + j, value=v)
        c.font = font(BLUE)
        c.number_format = '0.00"x"'
        c.fill = YELLOW
    A["scen_rows"].append(r)
a.cell(row=A["scen_rows"][0], column=7, value="Multipliers applied to the base partner ramp and to every price").font = font(italic=True, color="555555")

r += 2
section(a, r, "Partner network (clinics sharing data, end of year)", 7)
r += 1
header(a, r, ["Item"] + YEARS + ["Note"])
r += 1
inp_years(r, "Partner clinics, base case", [25, 120, 400, 900, 1500], NUM, "clinics", "Today's simulated network = 9 clinics. 1,500 clinics is about 4% of US practices", key=True)
r += 2
section(a, r, "Revenue assumptions (illustrative)", 7)
r += 1
header(a, r, ["Item", "Value", "", "", "", "", "Basis"])
r += 1
inp(r, "Premium scorecard price ($ per clinic per year)", 1800, USD, "sc_price", "Embedded analytics: group benchmarking beyond the free partner scorecard", key=True)
r += 1
inp(r, "Share of partner clinics buying premium scorecards", 0.40, PCT, "sc_adopt", "Assumption")
r += 1
inp(r, "Research report average price ($)", 40000, USD, "rep_price", "Research & insights (e.g., oral health report)")
r += 1
inp(r, "Data-license annual contract value ($)", 120000, USD, "lic_acv", "Data products: de-identified cohort, Health Index, resistance feed", key=True)
r += 1
inp(r, "Minimum partner clinics before data licenses are sold", 100, NUM, "lic_gate", "Privacy: at 9 clinics 8% of pets cannot be released; more partners shrink suppression")
r += 1
inp(r, "Scoring/API contract value ($ per client per year)", 80000, USD, "api_price", "Commercial models: drug-signal, demand-forecast scoring")
r += 1
inp(r, "Signal-validation study price ($ per study)", 250000, USD, "study_price", "EHR cohort study for a manufacturer or regulator", key=True)
r += 1
inp(r, "Clinics needed for common side-effect studies", 720, NUM, "gate_common", "From Model A2: ~80x today's 9-clinic network (docs/model_a2_results.md)")
r += 1
inp(r, "Clinics needed for rare side-effect studies", 1170, NUM, "gate_rare", "From Model A2: ~130x today's 9-clinic network")
r += 2
header(a, r, ["Volume by year"] + YEARS + ["Note"])
r += 1
inp_years(r, "Research reports sold", [2, 4, 6, 8, 10], NUM, "reports")
r += 1
inp_years(r, "Data licensees", [0, 2, 4, 6, 8], NUM, "licensees", "Sold only once the license gate is met")
r += 1
inp_years(r, "Scoring/API clients", [0, 1, 3, 5, 7], NUM, "api_clients")
r += 1
inp_years(r, "Common side-effect studies (if gate met)", [0, 1, 2, 2, 3], NUM, "studies_common")
r += 1
inp_years(r, "Rare side-effect studies (if gate met)", [0, 0, 1, 1, 2], NUM, "studies_rare")
r += 2
section(a, r, "Cost assumptions (illustrative)", 7)
r += 1
header(a, r, ["Item", "Value", "", "", "", "", "Basis"])
r += 1
inp(r, "Partner data incentive ($ per clinic per year)", 1000, USD, "incentive", "Free scorecards, integration support or rebates", key=True)
r += 1
inp(r, "Loaded cost per employee ($ per year)", 200000, USD, "loaded")
r += 1
inp(r, "Infrastructure, fixed ($ per year)", 60000, USD, "infra_fixed")
r += 1
inp(r, "Infrastructure per partner clinic ($ per year)", 50, USD, "infra_clinic")
r += 1
inp(r, "Study delivery cost (% of study revenue)", 0.40, PCT, "study_cogs", "Clinical review, outcome annotation, reporting")
r += 2
header(a, r, ["Headcount by year"] + YEARS + ["Note"])
r += 1
inp_years(r, "Team headcount (data science, engineering, partnerships, privacy)", [4, 7, 10, 13, 15], NUM, "headcount", key=True)

# ----------------------------------------------------------------------------- Model (three scenario blocks)
m = wb.create_sheet("Model")
m.column_dimensions["A"].width = 48
for j in range(2, 8):
    m.column_dimensions[L(j)].width = 15
m.cell(row=1, column=1, value="Five-year model by scenario ($)").font = font(bold=True, size=14)
m.cell(row=2, column=1, value="All cells are formulas driven by the Assumptions sheet.").font = font(italic=True, color="555555")
YC = [L(c) for c in range(2, 7)]  # model year columns B..F
AY = [L(c) for c in range(2, 7)]  # assumptions year columns B..F
blocks = {}
r = 4
for si, (sname, _, _) in enumerate(scen):
    sr = A["scen_rows"][si]
    ramp, price = f"Assumptions!$B${sr}", f"Assumptions!$C${sr}"
    section(m, r, f"{sname} scenario", 6)
    r += 1
    header(m, r, ["Line"] + YEARS)
    r += 1
    rows = {}

    def line(label, formulas, fmt=USD, bold=False, name=None):
        global r
        m.cell(row=r, column=1, value=label).font = font(bold=bold)
        for j, f in enumerate(formulas):
            c = m.cell(row=r, column=2 + j, value=f)
            c.number_format = fmt
            c.font = font(GREEN if "Assumptions!" in f and not any(op in f for op in "*+-/") else BLACK, bold=bold)
            if bold:
                c.border = Border(top=thin)
        if name:
            rows[name] = r
        r += 1

    line("Partner clinics", [f"=ROUND(Assumptions!{AY[j]}${A['clinics']}*{ramp},0)" for j in range(5)], NUM, name="clinics")
    cl = lambda j: f"{YC[j]}{rows['clinics']}"
    line("Embedded analytics: premium scorecards",
         [f"={cl(j)}*{A['sc_adopt']}*{A['sc_price']}*{price}" for j in range(5)], name="rev_sc")
    line("Research and insights: reports",
         [f"=Assumptions!{AY[j]}${A['reports']}*{A['rep_price']}*{price}" for j in range(5)], name="rev_rep")
    line("Data products: licenses",
         [f"=IF({cl(j)}>={A['lic_gate']},Assumptions!{AY[j]}${A['licensees']}*{A['lic_acv']}*{price},0)" for j in range(5)],
         name="rev_lic")
    line("Commercial models: scoring/API contracts",
         [f"=Assumptions!{AY[j]}${A['api_clients']}*{A['api_price']}*{price}" for j in range(5)], name="rev_api")
    line("Signal-validation studies, common side effects",
         [f"=IF({cl(j)}>={A['gate_common']},Assumptions!{AY[j]}${A['studies_common']}*{A['study_price']}*{price},0)"
          for j in range(5)], name="rev_sc_common")
    line("Signal-validation studies, rare side effects",
         [f"=IF({cl(j)}>={A['gate_rare']},Assumptions!{AY[j]}${A['studies_rare']}*{A['study_price']}*{price},0)"
          for j in range(5)], name="rev_sc_rare")
    line("Total revenue", [f"=SUM({YC[j]}{rows['rev_sc']}:{YC[j]}{rows['rev_sc_rare']})" for j in range(5)], bold=True,
         name="rev")
    line("Partner data incentives", [f"={cl(j)}*{A['incentive']}" for j in range(5)], name="c_inc")
    line("Team", [f"=Assumptions!{AY[j]}${A['headcount']}*{A['loaded']}" for j in range(5)], name="c_team")
    line("Infrastructure", [f"={A['infra_fixed']}+{cl(j)}*{A['infra_clinic']}" for j in range(5)], name="c_infra")
    line("Study delivery",
         [f"=({YC[j]}{rows['rev_sc_common']}+{YC[j]}{rows['rev_sc_rare']})*{A['study_cogs']}" for j in range(5)],
         name="c_study")
    line("Total costs", [f"=SUM({YC[j]}{rows['c_inc']}:{YC[j]}{rows['c_study']})" for j in range(5)], bold=True,
         name="cost")
    line("EBITDA", [f"={YC[j]}{rows['rev']}-{YC[j]}{rows['cost']}" for j in range(5)], bold=True, name="ebitda")
    line("EBITDA margin", [f"=IF({YC[j]}{rows['rev']}=0,0,{YC[j]}{rows['ebitda']}/{YC[j]}{rows['rev']})"
                           for j in range(5)], PCT, name="margin")
    line("Cumulative EBITDA", [f"={YC[0]}{rows['ebitda']}"] +
         [f"={YC[j - 1]}{r}+{YC[j]}{rows['ebitda']}" for j in range(1, 5)], name="cum")
    line("Helper: year number if EBITDA positive (else 99)",
         [f"=IF({YC[j]}{rows['ebitda']}>0,{j + 1},99)" for j in range(5)], NUM, name="be")
    blocks[sname] = rows
    r += 1

# ----------------------------------------------------------------------------- Summary
s = wb.create_sheet("Summary", 1)
s.column_dimensions["A"].width = 46
for col in "BCD":
    s.column_dimensions[col].width = 18
s.cell(row=1, column=1, value="Summary by scenario").font = font(bold=True, size=14)
s.cell(row=2, column=1, value="Illustrative assumptions; see Read Me.").font = font(italic=True, color="C00000")
header(s, 4, ["Measure", "Low", "Base", "High"])
meas = [("Partner clinics, Year 5", "clinics", "F", NUM), ("Total revenue, Year 5", "rev", "F", USD),
        ("EBITDA, Year 5", "ebitda", "F", USD), ("EBITDA margin, Year 5", "margin", "F", PCT),
        ("Cumulative EBITDA, Years 1-5", "cum", "F", USD)]
for i, (lbl, key, col, fmt) in enumerate(meas):
    rr = 5 + i
    s.cell(row=rr, column=1, value=lbl).font = font()
    for j, sname in enumerate(["Low", "Base", "High"]):
        c = s.cell(row=rr, column=2 + j, value=f"=Model!{col}{blocks[sname][key]}")
        c.font = font(GREEN)
        c.number_format = fmt
rr = 5 + len(meas)
s.cell(row=rr, column=1, value="First year with positive EBITDA").font = font()
for j, sname in enumerate(["Low", "Base", "High"]):
    b = blocks[sname]["be"]
    c = s.cell(row=rr, column=2 + j, value=f'=IF(MIN(Model!B{b}:F{b})=99,"Not within 5 years","Year "&MIN(Model!B{b}:F{b}))')
    c.font = font()
    c.alignment = Alignment(horizontal="right")
rr += 1
s.cell(row=rr, column=1, value="Revenue per partner clinic, Year 5").font = font()
for j, sname in enumerate(["Low", "Base", "High"]):
    b = blocks[sname]
    c = s.cell(row=rr, column=2 + j, value=f"=IF(Model!F{b['clinics']}=0,0,Model!F{b['rev']}/Model!F{b['clinics']})")
    c.number_format = USD
    c.font = font()
rr += 1
s.cell(row=rr, column=1, value="Partner incentive per clinic (cost)").font = font()
for j in range(3):
    c = s.cell(row=rr, column=2 + j, value=f"={A['incentive']}")
    c.number_format = USD
    c.font = font(GREEN)
rr += 1
s.cell(row=rr, column=1, value="Extra data licenses needed to break even in Year 5").font = font()
for j, sname in enumerate(["Low", "Base", "High"]):
    b = blocks[sname]
    sr = A["scen_rows"][j]
    c = s.cell(row=rr, column=2 + j,
               value=f"=ROUNDUP(MAX(0,-Model!F{b['ebitda']})/({A['lic_acv']}*Assumptions!$C${sr}),0)")
    c.number_format = NUM
    c.font = font()
rr += 2
header(s, rr, ["Revenue mix, Year 5 (Base)", "Revenue ($)", "Share"])
mix = [("Embedded analytics", "rev_sc"), ("Research and insights", "rev_rep"), ("Data products", "rev_lic"),
       ("Commercial models", "rev_api"), ("Validation studies, common", "rev_sc_common"),
       ("Validation studies, rare", "rev_sc_rare")]
first = rr + 1
for i, (lbl, key) in enumerate(mix):
    q = first + i
    s.cell(row=q, column=1, value=lbl).font = font()
    c = s.cell(row=q, column=2, value=f"=Model!F{blocks['Base'][key]}")
    c.font = font(GREEN)
    c.number_format = USD
    c = s.cell(row=q, column=3, value=f"=IF(SUM($B${first}:$B${first + len(mix) - 1})=0,0,B{q}/SUM($B${first}:$B${first + len(mix) - 1}))")
    c.number_format = PCT
    c.font = font()

# ----------------------------------------------------------------------------- Scale gates
g = wb.create_sheet("Scale Gates", 2)
g.column_dimensions["A"].width = 46
g.column_dimensions["B"].width = 18
for col in "CDE":
    g.column_dimensions[col].width = 16
g.column_dimensions["F"].width = 60
g.cell(row=1, column=1, value="When each product unlocks").font = font(bold=True, size=14)
header(g, 3, ["Capability", "Clinics needed", "Low", "Base", "High", "Evidence"])
gates = [("Clinic complication, dental and stewardship scorecards", 9, "Works at today's 9 clinics (docs/clinic_benchmarks_results.md)"),
         ("Data licenses (privacy-safe granularity)", "lic_gate", "Assumption; suppression 8% at 9 clinics"),
         ("Common side-effect validation studies", "gate_common", "~80x today's network (docs/model_a2_results.md)"),
         ("Rare side-effect validation studies", "gate_rare", "~130x today's network (docs/model_a2_results.md)")]
for i, (lbl, need, ev) in enumerate(gates):
    q = 4 + i
    g.cell(row=q, column=1, value=lbl).font = font()
    if isinstance(need, int):
        c = g.cell(row=q, column=2, value=need)
        c.font = font(BLUE)
    else:
        c = g.cell(row=q, column=2, value=f"={A[need]}")
        c.font = font(GREEN)
    c.number_format = NUM
    for j, sname in enumerate(["Low", "Base", "High"]):
        cr = blocks[sname]["clinics"]
        f = ("=IF(Model!B{r}>=$B{q},\"Year 1\",IF(Model!C{r}>=$B{q},\"Year 2\",IF(Model!D{r}>=$B{q},\"Year 3\","
             "IF(Model!E{r}>=$B{q},\"Year 4\",IF(Model!F{r}>=$B{q},\"Year 5\",\"Later\")))))").format(r=cr, q=q)
        cc = g.cell(row=q, column=3 + j, value=f)
        cc.font = font()
        cc.alignment = Alignment(horizontal="center")
    g.cell(row=q, column=6, value=ev).font = font(italic=True, color="555555")

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
